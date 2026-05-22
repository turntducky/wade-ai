import os
import re
import time
import json
import zlib
import queue
import base64
import sqlite3
import hashlib
import warnings
import threading

warnings.filterwarnings("ignore", category=SyntaxWarning)

try:
    import onnxruntime as _ort
    _ort.set_default_logger_severity(3)
except Exception:
    pass

import chromadb

from pathlib import Path
from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler
from typing import cast, Dict, Set, List, Optional, Tuple
from concurrent.futures import ThreadPoolExecutor, as_completed

from app.core.config import ConfigManager
from app.core.chroma_utils import UniversalEmbeddingFunction
from app.skills.indexing.code_chunker import LogicalCodeChunker

WADE_DIR = Path.home() / ".wade"
CHROMA_DB_DIR = WADE_DIR / "vector_store"
STATE_DB_PATH = WADE_DIR / "indexer_state.db"

WADE_DIR.mkdir(parents=True, exist_ok=True)
CHROMA_DB_DIR.mkdir(parents=True, exist_ok=True)

MAX_FILE_SIZE = 5 * 1024 * 1024  
MAX_WORKERS = min(8, os.cpu_count() or 4)
LIVE_WORKERS = max(1, MAX_WORKERS // 2)
CHROMA_BATCH_SIZE = 64
MAX_CHUNK_BATCH = 2000

ALLOWED_EXTENSIONS: Set[str] = {
    '.txt', '.md', '.rst', '.csv',
    '.pdf', '.docx', '.xlsx', '.pptx', '.odt', '.ods', '.odp',
    '.html', '.htm', '.css', '.scss', '.sass', '.less',
    '.js', '.ts', '.jsx', '.tsx', '.vue', '.svelte',
    '.py', '.pyw', '.pyx',
    '.c', '.cpp', '.cc', '.cxx', '.h', '.hpp', '.hxx', 
    '.cs', '.java',
    '.mq4', '.mq5', '.mqh',
    '.rs', '.go', '.rb', '.php', '.swift', '.kt', '.scala',
    '.sh', '.bash', '.ps1', '.bat', '.cmd',
    '.sql', '.graphql', '.gql',
    '.json', '.yaml', '.yml', '.toml', '.ini', '.xml',
    '.cfg', '.conf', '.env', '.uproject', '.uplugin'
}

ALLOWED_FILENAMES: Set[str] = {
    'dockerfile', 'docker-compose.yml', 'makefile', '.env', '.gitignore'
}

BLACKLIST_FILENAMES: Set[str] = {
    'identity.md', 'soul.md', 'agents.md', 'tools.md',
    'user.md', 'heartbeat.md', 'bootstrap.md', 'memory.md',
    'business.md', 'projects.md',
}

IGNORE_DIRS: Set[str] = {
    'node_modules', '.git', '__pycache__', '.wade', 'venv', 'env', '.idea', '.vscode',
    'build', 'dist', 'out', 'target', 'Release', 'Debug',
    '$Recycle.Bin', 'Recycle Bin', '.Trash', '.Trash-1000'
}

BLACKLIST_DIR_PATHS: Set[Path] = {
    (Path.home() / ".wade" / "workspace" / "memory").resolve(),
}

_MEMORY_FILE_RE = re.compile(r'^\d{2}-\d{2}-\d{2}(_[a-zA-Z0-9]+)?\.md$')

CORE_ZONES: List[Path] = [
    Path.home() / ".wade" / "workspace",
    Path(__file__).resolve().parent.parent.parent.parent
]

SYSTEM_ZONES: List[Path] = [
    Path.home() / ".wade" / "workspace",
    Path.home() / "OneDrive" / "Documents",
    Path.home() / "OneDrive" / "Desktop",
]

def _load_enabled_zones() -> List[Path]:
    """Load zones and directories based on user configuration."""
    try:
        config = ConfigManager.get()
        indexer_cfg = config.get("indexer", {})
        enabled_zone_names = indexer_cfg.get("enabled_zones", ["core", "system", "projects"])
        
        zones: List[Path] = []
        seen: set[Path] = set()

        def add_zone(p: Path):
            resolved = p.expanduser().resolve()
            if resolved not in seen and resolved.exists() and resolved.is_dir():
                seen.add(resolved)
                zones.append(resolved)

        if "core" in enabled_zone_names:
            for z in CORE_ZONES: add_zone(z)
        
        if "system" in enabled_zone_names:
            for z in SYSTEM_ZONES: add_zone(z)

        if "projects" in enabled_zone_names:
            raw_dirs = indexer_cfg.get("project_dirs", [])
            for raw in raw_dirs: add_zone(Path(str(raw)))

        custom_dirs = indexer_cfg.get("custom_dirs", [])
        for raw in custom_dirs: add_zone(Path(str(raw)))

        return zones
    except Exception:
        return CORE_ZONES + SYSTEM_ZONES

TARGET_ZONES: List[Path] = _load_enabled_zones()

IS_BOOTSTRAPPING = True
chroma_write_lock = threading.Lock()
sqlite_write_lock = threading.Lock()
db_pool_local = threading.local()

chroma_client = chromadb.PersistentClient(path=str(CHROMA_DB_DIR))
universal_ef = UniversalEmbeddingFunction()

def _get_or_recreate_collection(name: str, embedding_function):
    try:
        col = chroma_client.get_or_create_collection(name=name, embedding_function=embedding_function)

        try:
            col.peek(limit=1)
        except Exception as e:
            if "dimension" in str(e):
                raise ValueError(f"Dimension mismatch detected: {e}")
                
        return col

    except (ValueError, Exception) as e:
        err_msg = str(e)
        if "Embedding function conflict" in err_msg or "dimension" in err_msg:
            print(f"⚠️ Re-indexing required for collection '{name}': {err_msg}")
            try:
                chroma_client.delete_collection(name=name)
            except Exception: pass
            
            try:
                conn = sqlite3.connect(STATE_DB_PATH)
                conn.execute("DELETE FROM files")
                conn.commit()
                conn.close()
            except Exception as db_err:
                print(f"⚠️ Failed to reset state DB: {db_err}")

            return chroma_client.create_collection(name=name, embedding_function=embedding_function)
        raise e

_core_collection = None
_system_collection = None

def get_core_collection():
    global _core_collection
    if _core_collection is None:
        _core_collection = _get_or_recreate_collection(
            name="wade_core_workspace", 
            embedding_function=cast(chromadb.EmbeddingFunction, universal_ef)
        )
    return _core_collection

def get_system_collection():
    global _system_collection
    if _system_collection is None:
        _system_collection = _get_or_recreate_collection(
            name="wade_system_awareness", 
            embedding_function=cast(chromadb.EmbeddingFunction, universal_ef)
        )
    return _system_collection

code_chunker = LogicalCodeChunker()

def _is_core_path(filepath: str) -> bool:
    for zone in CORE_ZONES:
        if str(zone.absolute()) in filepath:
            return True
    return False

def get_db():
    if not hasattr(db_pool_local, "conn"):
        db_pool_local.conn = sqlite3.connect(STATE_DB_PATH, timeout=15.0)
        db_pool_local.conn.execute("PRAGMA journal_mode = WAL;")
        db_pool_local.conn.execute("PRAGMA synchronous = NORMAL;")
    return db_pool_local.conn

def init_state_db():
    conn = sqlite3.connect(STATE_DB_PATH)
    conn.execute("PRAGMA journal_mode = WAL;")

    conn.execute("""
        CREATE TABLE IF NOT EXISTS files (
            filepath TEXT PRIMARY KEY,
            raw_hash TEXT,
            stat_sig TEXT,
            chunk_ids TEXT
        )
    """)

    conn.execute("""
        CREATE TABLE IF NOT EXISTS chunks (
            chunk_id    TEXT PRIMARY KEY,
            filepath    TEXT NOT NULL,
            chunk_hash  TEXT NOT NULL,
            chunk_index INTEGER NOT NULL
        )
    """)
    conn.execute("CREATE INDEX IF NOT EXISTS idx_chunks_filepath ON chunks(filepath)")

    cursor = conn.cursor()
    cursor.execute("PRAGMA table_info(files)")
    columns = [col[1] for col in cursor.fetchall()]

    if "embed_model_id" not in columns:
        print("🔧 Migrating indexer state database (adding embed_model_id)...")
        conn.execute("ALTER TABLE files ADD COLUMN embed_model_id TEXT")

    conn.commit()
    conn.close()

init_state_db()

def _compress_ids(ids: List[str]) -> str:
    if not ids: return ""
    compressed = zlib.compress(json.dumps(ids).encode('utf-8'))
    return base64.b64encode(compressed).decode('utf-8')

def _decompress_ids(b64_str: str) -> List[str]:
    if not b64_str: return []
    try:
        return json.loads(zlib.decompress(base64.b64decode(b64_str.encode('utf-8'))).decode('utf-8'))
    except Exception: return []

def _get_active_model_id() -> str:
    config = ConfigManager.get()
    emb_cfg = config.get("active_suite", {}).get("embedding")
    if isinstance(emb_cfg, dict):
        return emb_cfg.get("repo") or emb_cfg.get("filename", "unknown")
    return str(emb_cfg or "all-MiniLM-L6-v2")

def _is_valid_file(file_path: Path, stat_obj: os.stat_result) -> bool:
    if stat_obj.st_size > MAX_FILE_SIZE or stat_obj.st_size == 0:
        return False
    name = file_path.name
    if name.lower() in BLACKLIST_FILENAMES:
        return False

    try:
        resolved = file_path.resolve()
        if any(resolved.is_relative_to(d) for d in BLACKLIST_DIR_PATHS):
            return False
    except Exception:
        pass

    if _MEMORY_FILE_RE.match(name):
        return False
    if name.startswith('~$') or name == '.DS_Store' or name.startswith('.'):
        return False
    if name.lower() in ALLOWED_FILENAMES:
        return True
    return file_path.suffix.lower() in ALLOWED_EXTENSIONS

def _get_raw_hash(file_path: Path, stat_obj: os.stat_result) -> str:
    try:
        with open(file_path, 'rb') as f:
            head = f.read(65536)
            try:
                f.seek(-65536, 2)
                tail = f.read(65536)
            except OSError: tail = b""
            payload = head + tail + str(stat_obj.st_size).encode('utf-8')
            return hashlib.blake2b(payload).hexdigest()
    except Exception: return ""

def _extract_text(file_path: Path, path_str: str) -> str:
    ext = file_path.suffix.lower()
    try:
        if ext == '.pdf':
            import fitz
            with fitz.open(path_str) as doc:
                return "\n".join([str(page.get_text()) for page in doc])
        elif ext == '.docx':
            import docx
            doc = docx.Document(path_str)
            return "\n".join([str(p.text) for p in doc.paragraphs])
        else:
            return file_path.read_text(encoding="utf-8", errors="ignore")
    except Exception: return ""

def _stream_chunks(text: str, chunk_size: int = 400, overlap: int = 50):
    words = text.split()
    step = chunk_size - overlap
    for i in range(0, len(words), step):
        chunk = " ".join(words[i:i + chunk_size])
        if chunk: yield chunk

def _chunk_markdown(text: str, max_words: int = 400):
    import re
    header_re = re.compile(r'^#{1,3} ', re.MULTILINE)
    lines = text.splitlines()
    sections = []
    current_header = None
    current_lines = []
    for line in lines:
        if header_re.match(line):
            if current_lines or current_header is not None:
                sections.append((current_header, current_lines))
            current_header = line
            current_lines = []
        else:
            current_lines.append(line)
    sections.append((current_header, current_lines))
    for header, content_lines in sections:
        body = "\n".join(content_lines).strip()
        full = f"{header}\n{body}".strip() if header else body
        if not full: continue
        if len(full.split()) <= max_words: yield full
        else: yield from _stream_chunks(full, max_words, overlap=50)

def _hash_chunk(text: str) -> str:
    return hashlib.blake2b(text.encode('utf-8')).hexdigest()[:32]

def _generate_chunk_id(path_str: str, chunk_hash: str) -> str:
    return hashlib.sha1(f"{path_str}::{chunk_hash}".encode('utf-8')).hexdigest()

def _process_file_for_ingestion(file_path: Path, stat_obj: os.stat_result) -> Optional[Dict]:
    path_str = str(file_path.absolute()) 
    stat_sig = f"{stat_obj.st_size}_{stat_obj.st_mtime}"
    active_model_id = _get_active_model_id()
    
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT raw_hash, stat_sig, chunk_ids, embed_model_id FROM files WHERE filepath = ?", (path_str,))
    row = cursor.fetchone()

    cursor.execute("SELECT chunk_id, chunk_hash FROM chunks WHERE filepath = ?", (path_str,))
    existing_chunk_map: Dict[str, str] = {r[0]: r[1] for r in cursor.fetchall()}
    existing_ids_set = set(existing_chunk_map.keys())
    chunk_table_populated = bool(existing_ids_set)

    model_changed = not row or row[3] != active_model_id

    if not model_changed and row and row[1] == stat_sig and chunk_table_populated:
        return {"action": "skip"}

    raw_hash = _get_raw_hash(file_path, stat_obj)
    if not raw_hash: return {"action": "skip"}

    if not model_changed and row and row[0] == raw_hash and chunk_table_populated:
        return {"action": "update_stat", "filepath": path_str, "raw_hash": raw_hash, "stat_sig": stat_sig, "chunk_ids": row[2], "model_id": active_model_id}

    content = _extract_text(file_path, path_str)
    if not content.strip(): return {"action": "skip"}

    modified_time = int(stat_obj.st_mtime)

    if code_chunker.is_supported(file_path):
        chunk_iter = code_chunker.stream_code_chunks(content, path_str)
    elif file_path.suffix.lower() == ".md":
        chunk_iter = _chunk_markdown(content)
    else:
        chunk_iter = _stream_chunks(content)

    all_chunks_data: List[tuple] = []
    seen_ids: Set[str] = set()
    for i, chunk in enumerate(chunk_iter):
        ch = _hash_chunk(chunk)
        cid = _generate_chunk_id(path_str, ch)
        if cid not in seen_ids:
            all_chunks_data.append((chunk, cid, ch, i))
            seen_ids.add(cid)

    if not all_chunks_data: return {"action": "skip"}

    new_ids_set = {d[1] for d in all_chunks_data}
    chunks_to_embed = [(text, cid, ch, idx) for text, cid, ch, idx in all_chunks_data if cid not in existing_ids_set]
    stale_ids = list(existing_ids_set - new_ids_set)

    if not chunks_to_embed and not stale_ids:
        return {
            "action": "update_stat",
            "filepath": path_str,
            "raw_hash": raw_hash,
            "stat_sig": stat_sig,
            "chunk_ids": _compress_ids(list(new_ids_set)),
            "model_id": active_model_id,
        }

    embed_chunks = [d[0] for d in chunks_to_embed]
    embed_ids    = [d[1] for d in chunks_to_embed]
    embed_metas  = [
        {"source": path_str, "filename": file_path.name, "chunk_index": d[3], "last_modified": modified_time}
        for d in chunks_to_embed
    ]
    new_chunk_hashes = {d[1]: d[2] for d in chunks_to_embed}

    return {
        "action": "ingest",
        "filepath": path_str,
        "filename": file_path.name,
        "raw_hash": raw_hash,
        "stat_sig": stat_sig,
        "chunks": embed_chunks,
        "ids": embed_ids,
        "all_ids": list(new_ids_set), 
        "old_ids": stale_ids,
        "chunk_count": len(embed_chunks),
        "metas": embed_metas,
        "model_id": active_model_id,
        "chunk_hashes": new_chunk_hashes,
    }

def _commit_batch(batch: List[Dict]):
    if not batch: return
    ingest_items = [item for item in batch if item["action"] == "ingest"]
    stat_updates = [item for item in batch if item["action"] == "update_stat"]
    core_ingest = [item for item in ingest_items if _is_core_path(item["filepath"])]
    system_ingest = [item for item in ingest_items if not _is_core_path(item["filepath"])]

    def _upsert_to_collection(items, collection):
        if not items: return
        all_chunks, all_metas, all_ids = [], [], []
        for item in items:
            all_chunks.extend(item["chunks"])
            all_ids.extend(item["ids"])
            all_metas.extend(item["metas"])
        
        with chroma_write_lock:
            try:
                for i in range(0, len(all_ids), CHROMA_BATCH_SIZE):
                    collection.upsert(
                        documents=all_chunks[i:i+CHROMA_BATCH_SIZE], 
                        metadatas=all_metas[i:i+CHROMA_BATCH_SIZE], 
                        ids=all_ids[i:i+CHROMA_BATCH_SIZE]
                    )
                    time.sleep(0.005) 
            except Exception as e:
                print(f" [!] Chroma Ingest Error: {e}")

    _upsert_to_collection(core_ingest, get_core_collection())
    _upsert_to_collection(system_ingest, get_system_collection())

    sqlite_files_data = []
    for item in ingest_items:
        all_ids = item.get("all_ids", item["ids"])
        sqlite_files_data.append((item["filepath"], item["raw_hash"], item["stat_sig"], _compress_ids(all_ids), item["model_id"]))
    for item in stat_updates:
        sqlite_files_data.append((item["filepath"], item["raw_hash"], item["stat_sig"], item["chunk_ids"], item["model_id"]))

    chunk_hash_inserts: List[tuple] = []
    stale_chunk_ids: List[str] = []
    for item in ingest_items:
        filepath = item["filepath"]
        id_to_index = {cid: item["metas"][i]["chunk_index"] for i, cid in enumerate(item["ids"])}
        for cid, ch in item.get("chunk_hashes", {}).items():
            chunk_hash_inserts.append((cid, filepath, ch, id_to_index.get(cid, 0)))
        stale_chunk_ids.extend(item.get("old_ids", []))

    with sqlite_write_lock:
        try:
            conn = get_db()
            if sqlite_files_data:
                conn.executemany("REPLACE INTO files (filepath, raw_hash, stat_sig, chunk_ids, embed_model_id) VALUES (?, ?, ?, ?, ?)", sqlite_files_data)
            if chunk_hash_inserts:
                conn.executemany("INSERT OR REPLACE INTO chunks (chunk_id, filepath, chunk_hash, chunk_index) VALUES (?, ?, ?, ?)", chunk_hash_inserts)
            for i in range(0, len(stale_chunk_ids), 500):
                chunk_batch = stale_chunk_ids[i:i + 500]
                ph = ",".join("?" * len(chunk_batch))
                conn.execute(f"DELETE FROM chunks WHERE chunk_id IN ({ph})", chunk_batch)
            conn.commit()
        except Exception as e:
            print(f" [!] DB Commit Error: {e}")

    all_old_ids_core = [id for item in core_ingest for id in item["old_ids"]]
    all_old_ids_system = [id for item in system_ingest for id in item["old_ids"]]
    
    for ids_to_del, coll in [(all_old_ids_core, get_core_collection()), (all_old_ids_system, get_system_collection())]:
        if ids_to_del:
            with chroma_write_lock:
                try:
                    for i in range(0, len(ids_to_del), CHROMA_BATCH_SIZE):
                        coll.delete(ids=ids_to_del[i:i+CHROMA_BATCH_SIZE])
                        time.sleep(0.005)
                except Exception: pass

def _purge_files_completely(filepaths: List[str]):
    if not filepaths: return
    with sqlite_write_lock:
        conn = get_db()
        cursor = conn.cursor()
        core_delete_ids, system_delete_ids = [], []
        for i in range(0, len(filepaths), 500):
            chunked_paths = filepaths[i:i+500]
            placeholders = ",".join("?" * len(chunked_paths))
            cursor.execute(f"SELECT filepath, chunk_ids FROM files WHERE filepath IN ({placeholders})", chunked_paths)
            for row in cursor.fetchall():
                filepath, chunk_ids = row[0], row[1]
                if chunk_ids:
                    ids_list = _decompress_ids(chunk_ids)
                    if _is_core_path(filepath): core_delete_ids.extend(ids_list)
                    else: system_delete_ids.extend(ids_list)
                    for j in range(0, len(ids_list), 500):
                        b = ids_list[j:j + 500]
                        ph2 = ",".join("?" * len(b))
                        cursor.execute(f"DELETE FROM chunks WHERE chunk_id IN ({ph2})", b)
            cursor.execute(f"DELETE FROM files WHERE filepath IN ({placeholders})", chunked_paths)
        conn.commit()
    with chroma_write_lock:
        for ids_to_delete, collection in [(core_delete_ids, get_core_collection()), (system_delete_ids, get_system_collection())]:
            if ids_to_delete:
                for i in range(0, len(ids_to_delete), CHROMA_BATCH_SIZE):
                    try:
                        collection.delete(ids=ids_to_delete[i:i+CHROMA_BATCH_SIZE])
                        time.sleep(0.005)
                    except Exception: pass

def bootstrap_differential_sync():
    global IS_BOOTSTRAPPING
    IS_BOOTSTRAPPING = True
    print("\n🔍 W.A.D.E. Initiating Dual-Track Sync...")
    
    get_core_collection()
    get_system_collection()
    
    disk_files: Set[str] = set()
    files_to_process: List[Tuple[Path, os.stat_result]] = []
    processed_paths: Set[str] = set()

    for zone in TARGET_ZONES:
        if not zone.exists(): continue
        for root, dirs, files in os.walk(zone):
            root_path = Path(root).resolve()
            dirs[:] = [
                d for d in dirs
                if not d.startswith('.')
                and d not in IGNORE_DIRS
                and (root_path / d).resolve() not in BLACKLIST_DIR_PATHS
            ]
            for file in files:
                file_path = Path(root) / file
                path_str = str(file_path.absolute())
                if path_str in processed_paths: continue
                try: stat_obj = file_path.stat()
                except OSError: continue
                if _is_valid_file(file_path, stat_obj):
                    disk_files.add(path_str)
                    files_to_process.append((file_path, stat_obj))
                    processed_paths.add(path_str)

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT filepath FROM files")
    existing_files = {row[0] for row in cursor.fetchall()}

    files_to_delete = list(existing_files - disk_files)
    if files_to_delete:
        print(f" 🧹 Purging {len(files_to_delete)} deleted files...")
        _purge_files_completely(files_to_delete)

    if files_to_process:
        total = len(files_to_process)
        print(f" 🔄 Scanning {total} files using {MAX_WORKERS} threads...")
        batch, chunk_count, done = [], 0, 0
        with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
            future_to_path = {executor.submit(_process_file_for_ingestion, path, stat): path for path, stat in files_to_process}
            for future in as_completed(future_to_path):
                done += 1
                if done % 100 == 0 or done == total:
                    print(f"\r  🧠 Ingesting: {done}/{total} files", end="", flush=True)
                try:
                    result = future.result()
                    if result and result["action"] in ("ingest", "update_stat"):
                        batch.append(result)
                        if result["action"] == "ingest": chunk_count += result["chunk_count"]
                        if chunk_count >= MAX_CHUNK_BATCH or len(batch) >= 500:
                            _commit_batch(batch)
                            batch.clear()
                            chunk_count = 0
                except Exception as e: print(f"\n [!] Ingestion Error: {e}")
        print()  
        if batch: _commit_batch(batch)

class QueuedLiveIndexHandler(FileSystemEventHandler):
    def __init__(self):
        self.last_processed = {} 
        self.event_queue = queue.Queue(maxsize=5000)        
        self.workers = []
        self._running = True
        for i in range(LIVE_WORKERS):
            t = threading.Thread(target=self._process_queue, daemon=True)
            t.start()
            self.workers.append(t)

    def _debounce(self, file_path: str) -> bool:
        if len(self.last_processed) > 5000: self.last_processed.clear()
        current_time = time.time()
        if file_path in self.last_processed and (current_time - self.last_processed[file_path]) < 0.1: return False
        self.last_processed[file_path] = current_time
        return True

    def _safe_put(self, action, filepath):
        if IS_BOOTSTRAPPING or not self._running: return  
        try: self.event_queue.put((action, filepath), block=False)
        except queue.Full: pass

    def on_created(self, event):
        if not event.is_directory:
            path_str = os.fsdecode(event.src_path)
            if self._debounce(path_str): self._safe_put("created", path_str)

    def on_modified(self, event):
        if not event.is_directory:
            path_str = os.fsdecode(event.src_path)
            if self._debounce(path_str): self._safe_put("modified", path_str)

    def on_deleted(self, event):
        if not event.is_directory:
            path_str = os.fsdecode(event.src_path)
            self._safe_put("deleted", path_str)

    def _process_queue(self):
        live_batch, filepaths_to_purge = [], []
        last_commit_time = time.time()
        while self._running:
            try:
                action, file_path = self.event_queue.get(timeout=1.0)
                path_obj = Path(file_path)
                path_str = str(path_obj.absolute())
                if action in ["created", "modified"]:
                    try:
                        stat_obj = path_obj.stat()
                        if _is_valid_file(path_obj, stat_obj):
                            result = _process_file_for_ingestion(path_obj, stat_obj)
                            if result and result["action"] in ("ingest", "update_stat"): live_batch.append(result)
                    except OSError: pass
                elif action == "deleted": filepaths_to_purge.append(path_str)
                self.event_queue.task_done()
            except queue.Empty: pass
            if live_batch and (len(live_batch) >= 20 or (time.time() - last_commit_time) > 1.0):
                _commit_batch(live_batch)
                live_batch.clear()
                last_commit_time = time.time()
            if filepaths_to_purge and (time.time() - last_commit_time) > 1.0:
                _purge_files_completely(filepaths_to_purge)
                filepaths_to_purge.clear()
                
    def stop(self): self._running = False

_observers = []
_live_handler = None

def _bootstrap_and_unblock():
    """Run differential sync in a background thread, then unlock the live watchdog."""
    global IS_BOOTSTRAPPING
    try:
        bootstrap_differential_sync()
    finally:
        IS_BOOTSTRAPPING = False
        print("\n[wade] Background indexing complete.")

def start_live_indexer():
    global _observers, _live_handler
    if _observers:
        return

    _live_handler = QueuedLiveIndexHandler()
    unique_zones = list({str(z.absolute()): z for z in TARGET_ZONES}.values())
    for zone in unique_zones:
        if zone.exists():
            obs = Observer()
            obs.schedule(_live_handler, str(zone), recursive=True)
            obs.start()
            _observers.append(obs)
    print(f"[wade] Indexer watching {len(_observers)} zone(s) — background sync starting...")
    threading.Thread(
        target=_bootstrap_and_unblock,
        daemon=True,
        name="wade-indexer-bootstrap",
    ).start()

def queue_file_for_index(filepath: str):
    if _live_handler and _live_handler._running: _live_handler._safe_put("modified", filepath)

def stop_live_indexer():
    global _observers, _live_handler
    print("\n🛑 Shutting down W.A.D.E. Indexer...")
    if _live_handler: _live_handler.stop()
    for obs in _observers:
        obs.stop()
        obs.join()
    _observers.clear()
    print("✅ Indexer cleanly stopped.")

if __name__ == "__main__": 
    start_live_indexer() 
    try: 
        while True: time.sleep(1)
    except KeyboardInterrupt: stop_live_indexer()